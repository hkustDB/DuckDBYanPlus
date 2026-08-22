#!/usr/bin/env python3
"""Plot the DuckDB 1.5 native, Yan, and Yan+ benchmark results.

The four plotted systems are:

* DuckDB1.5: native execution.
* DuckDB1.5-Yan (rewrite): the workbook's Yannakakis rewrite runtime.
* DuckDB1.5-Yan+ (rewrite): the workbook's Yannakakis+ rewrite runtime.
* DuckDB1.5-Yan+: the workbook's integrated Yannakakis+ runtime.

All plotted benchmark runtimes are read from the DuckDB 1.5 result workbook.
Each benchmark sheet uses the same schema: the first column contains query
names, followed by Origin, Yan, YanPlus_rewrite, and YanPlus runtimes in seconds.
The legacy CSV files are refresh inputs for the workbook, not figure inputs.
"""

from __future__ import annotations

import argparse
import math
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple


if "MPLCONFIGDIR" not in os.environ:
    _matplotlib_cache = Path(tempfile.gettempdir()) / "duckdb-v15-matplotlib"
    _matplotlib_cache.mkdir(parents=True, exist_ok=True)
    os.environ["MPLCONFIGDIR"] = str(_matplotlib_cache)
if "XDG_CACHE_HOME" not in os.environ:
    _font_cache = Path(tempfile.gettempdir()) / "duckdb-v15-cache"
    _font_cache.mkdir(parents=True, exist_ok=True)
    os.environ["XDG_CACHE_HOME"] = str(_font_cache)

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.ticker import FuncFormatter


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK_PATH = (
    Path.home()
    / "Library"
    / "CloudStorage"
    / "OneDrive-HKUSTConnect"
    / "DuckYan_1_5_results.xlsx"
)
DEFAULT_OUTPUT_DIR = REPOSITORY_ROOT / "figures" / "duckdb_v1_5"

DUCKDB = "DuckDB1.5"
YAN_REWRITE = "DuckDB1.5-Yan (rewrite)"
YANPLUS_REWRITE = r"DuckDB1.5-Yan$^{+}$ (rewrite)"
YANPLUS = r"DuckDB1.5-Yan$^{+}$"
SERIES_ORDER = (DUCKDB, YAN_REWRITE, YANPLUS_REWRITE, YANPLUS)
SERIES_STYLE = {
    DUCKDB: {"color": "#1f77b4", "hatch": "-"},
    YAN_REWRITE: {"color": "#D89000", "hatch": "//"},
    YANPLUS_REWRITE: {"color": "#30cf4d", "hatch": "\\"},
    YANPLUS: {"color": "#FFB6C1", "hatch": "||"},
}
BAR_EDGE_COLOR = "#505050"
BAR_LINEWIDTH = 0.3
TIME_LIMIT_SECONDS = 7200.0
TIMEOUT_LOG_CEILING_SECONDS = 1.0e4
JOB_Y_LIMITS = (1.0e-2, 1.0e4)
MERGED_FIGURE_SIZE = (50.0, 7.0)
MERGED_AXIS_LABEL_SIZE = 25
MERGED_TICK_LABEL_SIZE = 25
MERGED_SUBGRAPH_TITLE_SIZE = 25
MERGED_LEGEND_SIZE = 23.5
MERGED_ANNOTATION_SIZE = 12
JOB_AXIS_LABEL_SIZE = 22.5
JOB_TICK_LABEL_SIZE = 22
JOB_LEGEND_SIZE = 22.5
JOB_ANNOTATION_SIZE = 12
JOB_FIGURE_COUNT = 5
JOB_FIGURE_SIZE = (28.0, 5.0)
JOB_ALL_FIGURE_SIZE = JOB_FIGURE_SIZE
JOB_GROUP_WIDTH = 0.68
JOB_FAMILY_RANGES = ((1, 7), (8, 14), (15, 20), (21, 27), (28, 33))


@dataclass(frozen=True)
class FigureData:
    slug: str
    title: str
    query_ids: Sequence[str]
    query_labels: Sequence[str]
    series: Mapping[str, Sequence[Optional[float]]]
    notes: Sequence[str]


def standard_query(value: str, _path: Path) -> str:
    normalized = value.strip().lower()
    if re.fullmatch(r"q\d+", normalized):
        return normalized[1:]
    return normalized


def display_query(query: str) -> str:
    return query.upper() if query.startswith("q") else f"Q{query.upper()}"


WORKBOOK_COLUMN_SERIES = {
    "Origin": DUCKDB,
    "Yan": YAN_REWRITE,
    "YanPlus_rewrite": YANPLUS_REWRITE,
    "YanPlus": YANPLUS,
}


def workbook_sheet_rows(
    path: Path, sheet_name: str
) -> Tuple[List[str], Dict[str, List[Optional[float]]]]:
    """Read one benchmark sheet in row order from the authoritative workbook."""

    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "openpyxl is required to read the DuckDB 1.5 benchmark workbook"
        ) from error

    if not path.is_file():
        raise FileNotFoundError(f"missing benchmark workbook: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"{path} has no {sheet_name!r} worksheet")
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        workbook.close()

    if len(rows) < 2:
        raise ValueError(f"{path}/{sheet_name} has no benchmark data rows")
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    required = tuple(WORKBOOK_COLUMN_SERIES)
    missing = [column for column in required if column not in headers]
    if missing:
        raise ValueError(
            f"{path}/{sheet_name} is missing columns: {', '.join(missing)}"
        )

    query_ids: List[str] = []
    seen_queries = set()
    values_by_column: Dict[str, List[Optional[float]]] = {
        column: [] for column in required
    }
    for row_number, row in enumerate(rows[1:], start=2):
        raw_query = row[0] if row else None
        if raw_query is None or not str(raw_query).strip():
            continue
        query = standard_query(str(raw_query), path)
        if query in seen_queries:
            raise ValueError(
                f"duplicate query {raw_query!r} at {path}/{sheet_name}!A{row_number}"
            )
        seen_queries.add(query)
        query_ids.append(query)

        for column in required:
            column_index = headers.index(column)
            raw_value = row[column_index] if column_index < len(row) else None
            if raw_value is None or not str(raw_value).strip():
                values_by_column[column].append(None)
                continue
            try:
                value = float(raw_value)
            except (TypeError, ValueError) as error:
                raise ValueError(
                    f"invalid {column!r} value for {raw_query!r} at "
                    f"{path}/{sheet_name}!{row_number}"
                ) from error
            values_by_column[column].append(
                value if math.isfinite(value) and value > 0 else None
            )

    if not query_ids:
        raise ValueError(f"{path}/{sheet_name} has no named benchmark queries")
    return query_ids, values_by_column


def load_workbook_suite(
    workbook_path: Path,
    sheet_name: str,
    slug: str,
    title: str,
) -> FigureData:
    query_ids, values_by_column = workbook_sheet_rows(workbook_path, sheet_name)
    series = {
        WORKBOOK_COLUMN_SERIES[column]: values
        for column, values in values_by_column.items()
        if any(value is not None for value in values)
    }
    missing_cells = sum(
        value is None for values in values_by_column.values() for value in values
    )
    notes = [f"all runtimes loaded from workbook sheet {sheet_name!r}"]
    if missing_cells:
        notes.append(f"{missing_cells} unavailable workbook values are shown as N/A")
    return FigureData(
        slug=slug,
        title=title,
        query_ids=query_ids,
        query_labels=[display_query(query) for query in query_ids],
        series=series,
        notes=notes,
    )


def load_workbook_dsb(workbook_path: Path) -> FigureData:
    query_ids: List[str] = []
    query_labels: List[str] = []
    combined_series: Dict[str, List[Optional[float]]] = {
        series_name: [] for series_name in SERIES_ORDER
    }
    missing_cells = 0

    for sheet_name, prefix in (("dsbspj", "SPJ"), ("dsbagg", "AGG")):
        part_ids, values_by_column = workbook_sheet_rows(workbook_path, sheet_name)
        query_ids.extend(f"{prefix.lower()}-{query}" for query in part_ids)
        query_labels.extend(f"{prefix}-Q{query.upper()}" for query in part_ids)
        for column, series_name in WORKBOOK_COLUMN_SERIES.items():
            values = values_by_column[column]
            combined_series[series_name].extend(values)
            missing_cells += sum(value is None for value in values)

    series = {
        series_name: values
        for series_name, values in combined_series.items()
        if any(value is not None for value in values)
    }
    notes = ["all runtimes loaded from workbook sheets 'dsbspj' and 'dsbagg'"]
    if missing_cells:
        notes.append(f"{missing_cells} unavailable workbook values are shown as N/A")
    return FigureData(
        slug="dsb",
        title="DSB",
        query_ids=query_ids,
        query_labels=query_labels,
        series=series,
        notes=notes,
    )


def choose_scale(values: Iterable[float], requested: str) -> str:
    if requested != "auto":
        return requested
    positive = [value for value in values if value > 0]
    if not positive:
        return "linear"
    return "log" if max(positive) / min(positive) >= 50 else "linear"


def format_runtime_axis(axis, active_scale: str, values: Sequence[float]) -> None:
    axis.set_yscale(active_scale)
    if active_scale != "log":
        positive = [value for value in values if value > 0]
        if any(value >= TIME_LIMIT_SECONDS for value in positive):
            axis.set_ylim(0, TIME_LIMIT_SECONDS)
        else:
            axis.margins(y=0.05)
        return

    axis.yaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{value:.1e}"))
    positive = [value for value in values if value > 0]
    if positive:
        lower = 10 ** math.floor(math.log10(min(positive)))
        if any(value >= TIME_LIMIT_SECONDS for value in positive):
            upper = TIMEOUT_LOG_CEILING_SECONDS
        else:
            upper = 10 ** math.ceil(math.log10(max(positive)))
        if min(positive) / lower < 1.05:
            lower /= 10
        if (
            not any(value >= TIME_LIMIT_SECONDS for value in positive)
            and upper / max(positive) < 1.05
        ):
            upper *= 10
        axis.set_ylim(lower, upper)


def timeout_ceiling(active_scale: str, values: Sequence[float]) -> float:
    positive = [value for value in values if value > 0]
    if not positive:
        return TIME_LIMIT_SECONDS
    if active_scale != "log":
        return TIME_LIMIT_SECONDS if any(
            value >= TIME_LIMIT_SECONDS for value in positive
        ) else max(positive)

    if any(value >= TIME_LIMIT_SECONDS for value in positive):
        return TIMEOUT_LOG_CEILING_SECONDS

    upper = 10 ** math.ceil(math.log10(max(positive)))
    if upper / max(positive) < 1.05:
        upper *= 10
    return upper


def display_runtime(value: Optional[float], ceiling: float) -> Optional[float]:
    if value is not None and value >= TIME_LIMIT_SECONDS:
        return ceiling
    return value


def plot_panel(
    axis,
    data: FigureData,
    requested_scale: str,
    reference_values: Optional[Sequence[float]] = None,
    query_slots: Optional[int] = None,
    group_width: float = 0.88,
):
    series_names = [name for name in SERIES_ORDER if name in data.series]
    query_count = len(data.query_labels)
    if not 0 < group_width <= 1:
        raise ValueError("group_width must be greater than 0 and no greater than 1")
    bar_width = group_width / len(series_names)
    centers = list(range(query_count))
    all_values = [
        value
        for values in data.series.values()
        for value in values
        if value is not None and value > 0
    ]
    scale_values = list(reference_values) if reference_values is not None else all_values
    active_scale = choose_scale(scale_values, requested_scale)
    ceiling = timeout_ceiling(active_scale, scale_values)

    for series_index, series_name in enumerate(series_names):
        offset = -group_width / 2 + bar_width / 2 + series_index * bar_width
        positions = [center + offset for center in centers]
        values = data.series[series_name]
        display_values = [display_runtime(value, ceiling) for value in values]
        present_positions = [
            position
            for position, value in zip(positions, display_values)
            if value is not None
        ]
        present_values = [value for value in display_values if value is not None]
        style = SERIES_STYLE[series_name]
        axis.bar(
            present_positions,
            present_values,
            width=bar_width,
            color=style["color"],
            edgecolor=BAR_EDGE_COLOR,
            linewidth=BAR_LINEWIDTH,
            hatch=style["hatch"],
            zorder=3,
        )
        for position, value in zip(positions, values):
            if value is None:
                axis.annotate(
                    "N/A",
                    xy=(position, 0.01),
                    xycoords=("data", "axes fraction"),
                    rotation=90,
                    ha="center",
                    va="bottom",
                    fontsize=JOB_ANNOTATION_SIZE,
                    color="#555555",
                )

    axis.set_ylabel("Running Time (Sec)", fontsize=JOB_AXIS_LABEL_SIZE)
    axis.set_xticks(centers)
    axis.set_xticklabels(
        data.query_labels,
        rotation=0,
        ha="center",
        fontsize=JOB_TICK_LABEL_SIZE,
    )
    axis.tick_params(axis="y", labelsize=JOB_TICK_LABEL_SIZE)
    format_runtime_axis(axis, active_scale, scale_values)
    slot_count = query_count if query_slots is None else query_slots
    if slot_count < query_count:
        raise ValueError("query_slots cannot be smaller than the number of plotted queries")
    axis.set_xlim(-0.52, slot_count - 0.48)
    axis.grid(True, which="major", linestyle="--", linewidth=0.5, zorder=0)
    axis.set_axisbelow(True)
    return active_scale


def plot_combined_axis(axis, panels: Sequence[FigureData], requested_scale: str):
    all_values = [
        value
        for panel in panels
        for values in panel.series.values()
        for value in values
        if value is not None and value > 0
    ]
    active_scale = choose_scale(all_values, requested_scale)
    ceiling = timeout_ceiling(active_scale, all_values)
    query_labels: List[str] = []
    section_bounds = []
    cursor = 0

    for panel in panels:
        series_names = [name for name in SERIES_ORDER if name in panel.series]
        group_width = 0.88
        bar_width = group_width / len(series_names)
        centers = list(range(cursor, cursor + len(panel.query_ids)))
        section_bounds.append((centers[0], centers[-1], panel.title))
        query_labels.extend(panel.query_labels)

        for series_index, series_name in enumerate(series_names):
            offset = -group_width / 2 + bar_width / 2 + series_index * bar_width
            positions = [center + offset for center in centers]
            values = panel.series[series_name]
            display_values = [display_runtime(value, ceiling) for value in values]
            present_positions = [
                position
                for position, value in zip(positions, display_values)
                if value is not None
            ]
            present_values = [value for value in display_values if value is not None]
            style = SERIES_STYLE[series_name]
            axis.bar(
                present_positions,
                present_values,
                width=bar_width,
                color=style["color"],
                edgecolor=BAR_EDGE_COLOR,
                linewidth=BAR_LINEWIDTH,
                hatch=style["hatch"],
                zorder=3,
            )
            for position, value in zip(positions, values):
                if value is None:
                    axis.annotate(
                        "N/A",
                        xy=(position, 0.01),
                        xycoords=("data", "axes fraction"),
                        rotation=90,
                        ha="center",
                        va="bottom",
                        fontsize=MERGED_ANNOTATION_SIZE,
                        color="#555555",
                    )
        cursor += len(panel.query_ids)

    for first, last, title in section_bounds:
        midpoint = (first + last) / 2
        axis.text(
            midpoint,
            -0.075,
            title,
            transform=axis.get_xaxis_transform(),
            ha="center",
            va="top",
            fontsize=MERGED_SUBGRAPH_TITLE_SIZE,
            clip_on=False,
        )

    centers = list(range(len(query_labels)))
    axis.set_ylabel("Running Time (Sec)", fontsize=MERGED_AXIS_LABEL_SIZE)
    axis.set_xticks(centers)
    axis.set_xticklabels(
        query_labels,
        rotation=0,
        ha="center",
        fontsize=MERGED_TICK_LABEL_SIZE,
    )
    axis.tick_params(axis="x", pad=2)
    axis.tick_params(axis="y", labelsize=MERGED_TICK_LABEL_SIZE)
    format_runtime_axis(axis, active_scale, all_values)
    axis.set_xlim(-0.52, len(query_labels) - 0.48)
    axis.grid(True, which="major", linestyle="--", linewidth=0.5, zorder=0)
    axis.set_axisbelow(True)
    return active_scale


def legend_handles(series_names: Sequence[str]):
    return [
        Patch(
            facecolor=SERIES_STYLE[name]["color"],
            edgecolor=BAR_EDGE_COLOR,
            linewidth=BAR_LINEWIDTH,
            hatch=SERIES_STYLE[name]["hatch"],
            label=name,
        )
        for name in SERIES_ORDER
        if name in series_names
    ]


def save_figure(
    figure,
    output_dir: Path,
    basename: str,
    formats: Sequence[str],
    dpi: int,
    *,
    tight: bool = True,
):
    output_dir.mkdir(parents=True, exist_ok=True)
    destinations = []
    for extension in formats:
        destination = output_dir / f"{basename}.{extension}"
        figure.savefig(
            destination,
            dpi=dpi,
            bbox_inches="tight" if tight else None,
            facecolor="white",
        )
        destinations.append(destination)
    plt.close(figure)
    return destinations


def plot_merged(
    panels: Sequence[FigureData],
    basename: str,
    output_dir: Path,
    formats: Sequence[str],
    dpi: int,
    scale: str,
):
    figure, axis = plt.subplots(figsize=MERGED_FIGURE_SIZE)
    active_scale = plot_combined_axis(axis, panels, scale)

    present_series = {
        series_name for panel in panels for series_name in panel.series
    }
    axis.legend(
        handles=legend_handles(tuple(present_series)),
        loc="lower center",
        bbox_to_anchor=(0.5, 1.15),
        ncol=min(4, len(present_series)),
        frameon=False,
        fontsize=MERGED_LEGEND_SIZE,
        borderaxespad=0,
    )
    figure.subplots_adjust(top=0.83, bottom=0.19, left=0.04, right=0.99)
    return (
        save_figure(
            figure,
            output_dir,
            basename,
            formats,
            dpi,
            tight=False,
        ),
        active_scale,
    )


def subset_indices(data: FigureData, first_index: int, last_index: int) -> FigureData:
    indices = list(range(first_index, min(last_index, len(data.query_ids))))
    if not indices:
        raise ValueError(f"JOB partition {first_index + 1}-{last_index} is empty")
    return FigureData(
        slug=f"job_figure_{first_index + 1}_{min(last_index, len(data.query_ids))}",
        title=f"JOB queries {first_index + 1}-{min(last_index, len(data.query_ids))}",
        query_ids=[data.query_ids[index] for index in indices],
        query_labels=[data.query_labels[index] for index in indices],
        series={
            name: [values[index] for index in indices]
            for name, values in data.series.items()
        },
        notes=data.notes,
    )


def subset_families(
    data: FigureData, first_family: int, last_family: int
) -> FigureData:
    indices = []
    for index, query in enumerate(data.query_ids):
        match = re.match(r"(\d+)", query)
        if match and first_family <= int(match.group(1)) <= last_family:
            indices.append(index)
    if not indices:
        raise ValueError(f"JOB partition {first_family}-{last_family} is empty")
    return FigureData(
        slug=f"job_{first_family}_{last_family}",
        title=f"JOB queries {first_family}-{last_family}",
        query_ids=[data.query_ids[index] for index in indices],
        query_labels=[data.query_labels[index] for index in indices],
        series={
            name: [values[index] for index in indices]
            for name, values in data.series.items()
        },
        notes=data.notes,
    )


def plot_job_family_partitions(
    data: FigureData,
    output_dir: Path,
    formats: Sequence[str],
    dpi: int,
    scale: str,
):
    panels = [
        (first_family, last_family, subset_families(data, first_family, last_family))
        for first_family, last_family in JOB_FAMILY_RANGES
    ]
    reference_values = [
        value
        for values in data.series.values()
        for value in values
        if value is not None and value > 0
    ]
    results = []
    for partition_index, (first_family, last_family, panel) in enumerate(panels):
        figure, axis = plt.subplots(figsize=JOB_FIGURE_SIZE)
        active_scale = plot_panel(
            axis,
            panel,
            scale,
            reference_values=reference_values,
            group_width=JOB_GROUP_WIDTH,
        )
        axis.set_ylim(*JOB_Y_LIMITS)
        if partition_index == 0:
            axis.legend(
                handles=legend_handles(tuple(panel.series)),
                loc="lower center",
                bbox_to_anchor=(0.5, 1.01),
                ncol=4,
                frameon=False,
                fontsize=JOB_LEGEND_SIZE,
                borderaxespad=0,
            )
        figure.subplots_adjust(top=0.89, bottom=0.10, left=0.065, right=0.995)
        basename = f"job_benchmark_{first_family}_{last_family}"
        destinations = save_figure(
            figure,
            output_dir,
            basename,
            formats,
            dpi,
            tight=False,
        )
        results.append(
            (
                basename,
                destinations,
                active_scale,
                len(panel.query_ids),
                len(panel.query_ids),
            )
        )
    return results


def plot_job_partitions(
    data: FigureData,
    output_dir: Path,
    formats: Sequence[str],
    dpi: int,
    scale: str,
):
    results = []
    slots_per_figure = math.ceil(len(data.query_ids) / JOB_FIGURE_COUNT)
    reference_values = [
        value
        for values in data.series.values()
        for value in values
        if value is not None and value > 0
    ]
    for figure_number in range(1, JOB_FIGURE_COUNT + 1):
        first_index = (figure_number - 1) * slots_per_figure
        last_index = first_index + slots_per_figure
        panel = subset_indices(data, first_index, last_index)
        figure, axis = plt.subplots(figsize=JOB_FIGURE_SIZE)
        active_scale = plot_panel(
            axis,
            panel,
            scale,
            reference_values=reference_values,
            group_width=JOB_GROUP_WIDTH,
        )
        axis.set_ylim(*JOB_Y_LIMITS)
        if figure_number == 1:
            axis.legend(
                handles=legend_handles(tuple(panel.series)),
                loc="lower center",
                bbox_to_anchor=(0.5, 1.01),
                ncol=4,
                frameon=False,
                fontsize=JOB_LEGEND_SIZE,
                borderaxespad=0,
            )
        figure.subplots_adjust(top=0.89, bottom=0.10, left=0.065, right=0.995)
        basename = f"job_benchmark_figure_{figure_number}"
        destinations = save_figure(
            figure,
            output_dir,
            basename,
            formats,
            dpi,
            tight=False,
        )
        results.append(
            (
                basename,
                destinations,
                active_scale,
                len(panel.query_ids),
                len(panel.query_ids),
            )
        )
    return results


def plot_job_all(
    data: FigureData,
    output_dir: Path,
    formats: Sequence[str],
    dpi: int,
    scale: str,
):
    reference_values = [
        value
        for values in data.series.values()
        for value in values
        if value is not None and value > 0
    ]
    figure, axis = plt.subplots(figsize=JOB_ALL_FIGURE_SIZE)
    active_scale = plot_panel(
        axis,
        data,
        scale,
        reference_values=reference_values,
        group_width=JOB_GROUP_WIDTH,
    )
    axis.set_ylim(*JOB_Y_LIMITS)
    labeled_indices = list(range(0, len(data.query_labels), 5))
    axis.set_xticks(labeled_indices)
    axis.set_xticklabels(
        [data.query_labels[index] for index in labeled_indices],
        rotation=0,
        ha="center",
        fontsize=JOB_TICK_LABEL_SIZE,
    )
    axis.legend(
        handles=legend_handles(tuple(data.series)),
        loc="lower center",
        bbox_to_anchor=(0.5, 1.01),
        ncol=4,
        frameon=False,
        fontsize=JOB_LEGEND_SIZE,
        borderaxespad=0,
    )
    figure.subplots_adjust(top=0.89, bottom=0.10, left=0.065, right=0.995)
    destinations = save_figure(
        figure,
        output_dir,
        "job_benchmark_all",
        formats,
        dpi,
        tight=False,
    )
    return destinations, active_scale, len(data.query_ids)


def parse_args(argv: Optional[Sequence[str]] = None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"figure destination (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK_PATH,
        help=f"authoritative benchmark result workbook (default: {DEFAULT_WORKBOOK_PATH})",
    )
    parser.add_argument(
        "--benchmarks",
        nargs="+",
        choices=("all", "graph", "lsqb", "tpch", "job", "dsb"),
        default=("all",),
        help="benchmark suites to include",
    )
    parser.add_argument(
        "--formats",
        nargs="+",
        choices=("png", "pdf", "svg"),
        default=("pdf",),
        help="output formats",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=1200,
        help="rasterization resolution for saved figures (default: 1200 dpi)",
    )
    parser.add_argument(
        "--scale",
        choices=("auto", "linear", "log"),
        default="log",
        help="y-axis scale (default: log, matching the reference drawing style)",
    )
    return parser.parse_args(argv)


def report_notes(panels: Sequence[FigureData]):
    for panel in panels:
        for note in panel.notes:
            print(f"  {panel.title} note: {note}")


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    workbook_path = args.workbook.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    formats = tuple(dict.fromkeys(args.formats))
    selected = (
        {"graph", "lsqb", "tpch", "job", "dsb"}
        if "all" in args.benchmarks
        else set(args.benchmarks)
    )
    loaders = {
        "graph": lambda: load_workbook_suite(
            workbook_path, "graph", "graph", "Graph"
        ),
        "lsqb": lambda: load_workbook_suite(
            workbook_path, "lsqb", "lsqb", "LSQB"
        ),
        "tpch": lambda: load_workbook_suite(
            workbook_path, "tpch", "tpch", "TPC-H"
        ),
        "job": lambda: load_workbook_suite(
            workbook_path, "job", "job", "JOB"
        ),
        "dsb": lambda: load_workbook_dsb(workbook_path),
    }
    available = {name: loaders[name]() for name in selected}

    graph_lsqb = [available[name] for name in ("graph", "lsqb") if name in selected]
    if graph_lsqb:
        destinations, active_scale = plot_merged(
            graph_lsqb,
            "graph_lsqb_benchmark",
            output_dir,
            formats,
            args.dpi,
            args.scale,
        )
        print(
            "graph/lsqb: "
            + active_scale
            + " scale -> "
            + ", ".join(str(path) for path in destinations)
        )
        report_notes(graph_lsqb)

    tpch_dsb = [available[name] for name in ("tpch", "dsb") if name in selected]
    if tpch_dsb:
        destinations, active_scale = plot_merged(
            tpch_dsb,
            "tpch_dsb_benchmark",
            output_dir,
            formats,
            args.dpi,
            args.scale,
        )
        print(
            "tpch/dsb: "
            + active_scale
            + " scale -> "
            + ", ".join(str(path) for path in destinations)
        )
        report_notes(tpch_dsb)

    if "job" in selected:
        for basename, destinations, active_scale, count, slots in plot_job_family_partitions(
            available["job"], output_dir, formats, args.dpi, args.scale
        ):
            print(
                f"{basename}: {count} queries in {slots} plot slots, "
                f"{active_scale} scale -> "
                + ", ".join(str(path) for path in destinations)
            )
        for basename, destinations, active_scale, count, slots in plot_job_partitions(
            available["job"], output_dir, formats, args.dpi, args.scale
        ):
            print(
                f"{basename}: {count} queries in {slots} plot slots, "
                f"{active_scale} scale -> "
                + ", ".join(str(path) for path in destinations)
            )
        destinations, active_scale, count = plot_job_all(
            available["job"], output_dir, formats, args.dpi, args.scale
        )
        print(
            f"job_benchmark_all: {count} queries, {active_scale} scale -> "
            + ", ".join(str(path) for path in destinations)
        )
        report_notes((available["job"],))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
